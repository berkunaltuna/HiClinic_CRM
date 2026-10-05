from __future__ import annotations

import io
import threading
import uuid
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest
from sqlalchemy.orm import sessionmaker

from app.api.events import create_appointment as api_create_appointment
from app.db.models import Appointment, AuditLog, Customer, Deal, Event, ImportBatch, ImportRow, User
from app.schemas.event import AppointmentCreate
from app.services.event_import import run_commit

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EXPORT_HEADERS = [
    "Customer ID", "Booking ID", "Name", "Phone", "Email", "Treatment Interest",
    "Event", "Date", "Appointment Time", "Booking Status", "Notes",
]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def make_xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_event(client, headers, *, name="Test Event", day="2026-09-01", slot_capacity=1, slot_minutes=30, start="09:00", end="17:00"):
    r = client.post(
        "/events",
        json={
            "name": name,
            "location": "Clinic",
            "starts_on": day,
            "ends_on": day,
            "default_slot_minutes": slot_minutes,
            "slot_capacity": slot_capacity,
            "days": [{"day": day, "start_time": start, "end_time": end, "slot_minutes": slot_minutes}],
        },
        headers=headers,
    )
    assert r.status_code == 201, r.text
    return r.json()


def default_mapping() -> dict[str, str]:
    return {
        "customer_id": "Customer ID",
        "booking_id": "Booking ID",
        "name": "Name",
        "phone": "Phone",
        "email": "Email",
        "treatment_interest": "Treatment Interest",
        "event": "Event",
        "date": "Date",
        "appointment_time": "Appointment Time",
        "booking_status": "Booking Status",
        "notes": "Notes",
    }


def upload(client, headers, event_id, rows: list[list]):
    content = make_xlsx_bytes(EXPORT_HEADERS, rows)
    r = client.post(
        f"/events/{event_id}/imports/upload",
        headers=headers,
        files={"file": ("import.xlsx", content, XLSX_CT)},
    )
    assert r.status_code == 201, r.text
    return r.json()


def preview(client, headers, event_id, batch_id, mapping=None):
    r = client.post(
        f"/events/{event_id}/imports/{batch_id}/preview",
        headers=headers,
        json={"mapping": mapping or default_mapping()},
    )
    assert r.status_code == 200, r.text
    return r.json()


def commit(client, headers, event_id, batch_id):
    return client.post(f"/events/{event_id}/imports/{batch_id}/commit", headers=headers)


def row_for(name, phone="", email="", treatment="", event="Test Event", date="2026-09-01", time="09:00", status="booked", notes="", customer_id="", booking_id=""):
    return [customer_id, booking_id, name, phone, email, treatment, event, date, time, status, notes]


def register(client, email=None):
    email = email or f"user_{uuid.uuid4().hex[:10]}@example.com"
    r = client.post("/auth/register", json={"email": email, "password": "ChangeMe123!"})
    assert r.status_code == 201, r.text
    token = r.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}, email


def db_user(db, email) -> User:
    return db.query(User).filter(User.email == email).first()


# ---------------------------------------------------------------------------
# 1. Export round trip
# ---------------------------------------------------------------------------


def test_export_round_trip(client, auth_headers):
    event = create_event(client, auth_headers)
    r = client.post(
        "/customers",
        json={"name": "Jane Export", "email": "jane.export.test@example.com", "phone": "+447700900199"},
        headers=auth_headers,
    )
    customer_id = r.json()["id"]
    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": customer_id, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00", "notes": "hello"},
        headers=auth_headers,
    )
    assert r.status_code == 201, r.text

    r = client.get(f"/events/{event['id']}/export.xlsx", headers=auth_headers)
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header_row == EXPORT_HEADERS
    data_row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert data_row[2] == "Jane Export"
    assert data_row[7] == "2026-09-01"
    assert data_row[8] == "09:00"
    assert data_row[9] == "booked"


# ---------------------------------------------------------------------------
# 2-8: core matching/dedup behavior
# ---------------------------------------------------------------------------


def test_import_happy_path_creates_customer_and_booking(client, auth_headers, db):
    event = create_event(client, auth_headers)
    up = upload(client, auth_headers, event["id"], [row_for("Alice New", phone="07700900001", email="alice@example.com")])
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["batch"]["customers_created"] == 1
    assert prev["batch"]["bookings_created"] == 1
    res = commit(client, auth_headers, event["id"], up["batch_id"])
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["batch"]["customers_created"] == 1
    assert body["batch"]["bookings_created"] == 1
    cust = db.query(Customer).filter(Customer.name == "Alice New").first()
    assert cust is not None
    appt = db.query(Appointment).filter(Appointment.customer_id == cust.id).first()
    assert appt is not None
    assert appt.status == "booked"
    assert cust.stage == "appointment_booked"


def test_import_phone_match_updates_not_duplicates(client, auth_headers, db):
    event = create_event(client, auth_headers)
    r = client.post("/customers", json={"name": "Bob Existing", "phone": "07700 900002"}, headers=auth_headers)
    assert r.status_code == 201, r.text

    up = upload(client, auth_headers, event["id"], [row_for("Bob Existing", phone="+447700900002", email="bob@example.com")])
    preview(client, auth_headers, event["id"], up["batch_id"])
    res = commit(client, auth_headers, event["id"], up["batch_id"])
    body = res.json()
    assert body["batch"]["customers_created"] == 0
    assert body["batch"]["customers_updated"] == 1
    assert db.query(Customer).filter(Customer.name == "Bob Existing").count() == 1
    cust = db.query(Customer).filter(Customer.name == "Bob Existing").first()
    assert cust.email == "bob@example.com"


def test_import_case_insensitive_email_match(client, auth_headers, db):
    event = create_event(client, auth_headers)
    r = client.post("/customers", json={"name": "Cara Case", "email": "Cara@Example.com"}, headers=auth_headers)
    assert r.status_code == 201, r.text

    up = upload(client, auth_headers, event["id"], [row_for("Cara Case", email="cara@example.com")])
    preview(client, auth_headers, event["id"], up["batch_id"])
    res = commit(client, auth_headers, event["id"], up["batch_id"])
    body = res.json()
    assert body["batch"]["customers_created"] == 0
    assert db.query(Customer).filter(Customer.name == "Cara Case").count() == 1


def test_import_ambiguous_phone_and_email_point_to_different_customers_flagged(client, auth_headers, db):
    event = create_event(client, auth_headers)
    client.post("/customers", json={"name": "Customer A", "phone": "07700900010"}, headers=auth_headers)
    client.post("/customers", json={"name": "Customer B", "email": "b@example.com"}, headers=auth_headers)

    up = upload(client, auth_headers, event["id"], [row_for("Ambiguous", phone="+447700900010", email="b@example.com")])
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["rows"][0]["status"] == "flagged"
    before_customers = db.query(Customer).count()
    res = commit(client, auth_headers, event["id"], up["batch_id"])
    assert res.json()["batch"]["rows_flagged"] == 1
    assert db.query(Customer).count() == before_customers
    assert db.query(Appointment).filter(Appointment.event_id == uuid.UUID(event["id"])).count() == 0


def test_import_multiple_existing_customers_share_phone_flagged(client, auth_headers, db):
    event = create_event(client, auth_headers)
    client.post("/customers", json={"name": "Dup One", "phone": "07700900020"}, headers=auth_headers)
    client.post("/customers", json={"name": "Dup Two", "phone": "07700900020"}, headers=auth_headers)

    up = upload(client, auth_headers, event["id"], [row_for("Someone", phone="+447700900020")])
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["rows"][0]["status"] == "flagged"
    assert "multiple" in prev["rows"][0]["reasons"][0]


def test_import_no_stable_identifier_flagged(client, auth_headers, db):
    event = create_event(client, auth_headers)
    up = upload(client, auth_headers, event["id"], [row_for("No Identifier")])
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["rows"][0]["status"] == "flagged"
    assert db.query(Customer).filter(Customer.name == "No Identifier").count() == 0


def test_import_within_batch_duplicate_consolidation(client, auth_headers, db):
    # Same person, same exact slot (different rows in the sheet referring to the same booking) --
    # this exercises pure customer consolidation without tripping the separate "existing active
    # booking at a different time" review rule (covered by other tests), since that rule requires
    # an explicit booking_id to disambiguate a genuine reschedule from an accidental duplicate row.
    event = create_event(client, auth_headers, slot_capacity=5)
    rows = [
        row_for("Dup Row 1", phone="07700900030", time="09:00"),
        row_for("Dup Row 1 Again", phone="+447700900030", time="09:00"),
    ]
    up = upload(client, auth_headers, event["id"], rows)
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["batch"]["customers_created"] == 1
    assert prev["batch"]["duplicates_merged"] == 1
    assert prev["rows"][1]["status"] == "unchanged"
    res = commit(client, auth_headers, event["id"], up["batch_id"])
    body = res.json()
    assert body["batch"]["customers_created"] == 1
    assert body["batch"]["duplicates_merged"] == 1
    assert db.query(Customer).filter(Customer.phone == "+447700900030").count() == 1
    cust = db.query(Customer).filter(Customer.phone == "+447700900030").first()
    assert db.query(Appointment).filter(Appointment.customer_id == cust.id).count() == 1


# ---------------------------------------------------------------------------
# 9. Preview matches commit exactly (duplicates + capacity)
# ---------------------------------------------------------------------------


def test_preview_matches_commit_for_duplicates_and_capacity(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=1)
    rows = [
        row_for("Cap Row 1", phone="07700900040", time="09:00"),
        row_for("Cap Row 1 Dup", phone="+447700900040", time="09:00"),
        row_for("Cap Row 2", phone="07700900041", time="09:00"),
    ]
    up = upload(client, auth_headers, event["id"], rows)
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    preview_statuses = [r["status"] for r in prev["rows"]]

    res = commit(client, auth_headers, event["id"], up["batch_id"])
    body = res.json()
    commit_statuses = [r["status"] for r in body["rows"]]

    assert preview_statuses == commit_statuses
    assert preview_statuses[0] == "created"
    # second row is same customer, same exact slot -> idempotent "unchanged" update-in-place
    assert preview_statuses[1] == "unchanged"
    # third row is a different customer competing for the same (now full) slot -> flagged
    assert preview_statuses[2] == "flagged"
    assert prev["batch"]["bookings_created"] == body["batch"]["bookings_created"] == 1


# ---------------------------------------------------------------------------
# 10-13: booking resolution strictness
# ---------------------------------------------------------------------------


def test_import_booking_id_validation(client, auth_headers, db):
    event = create_event(client, auth_headers)
    other_event = create_event(client, auth_headers, name="Other Event", day="2026-09-02")

    r = client.post("/customers", json={"name": "Booking Owner", "phone": "07700900050"}, headers=auth_headers)
    customer_id = r.json()["id"]
    r = client.post(
        f"/events/{other_event['id']}/appointments",
        json={"customer_id": customer_id, "starts_at": "2026-09-02T09:00:00", "ends_at": "2026-09-02T09:30:00"},
        headers=auth_headers,
    )
    wrong_event_booking_id = r.json()["id"]

    rows = [
        row_for("Booking Owner", phone="07700900050", booking_id=str(uuid.uuid4())),  # doesn't exist
        row_for("Booking Owner", phone="07700900050", booking_id=wrong_event_booking_id),  # wrong event
    ]
    up = upload(client, auth_headers, event["id"], rows)
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["rows"][0]["status"] == "flagged"
    assert prev["rows"][1]["status"] == "flagged"
    assert db.query(Appointment).filter(Appointment.event_id == uuid.UUID(event["id"])).count() == 0


def test_import_multiple_exact_slot_matches_flagged(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=5)
    r = client.post("/customers", json={"name": "Double Booked", "phone": "07700900060"}, headers=auth_headers)
    customer_id = r.json()["id"]
    for _ in range(2):
        r = client.post(
            f"/events/{event['id']}/appointments",
            json={"customer_id": customer_id, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00"},
            headers=auth_headers,
        )
        assert r.status_code == 201, r.text

    up = upload(client, auth_headers, event["id"], [row_for("Double Booked", phone="07700900060", time="09:00")])
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["rows"][0]["status"] == "flagged"
    assert "multiple" in prev["rows"][0]["reasons"][0]


def test_import_reschedule_revalidates_capacity(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=1)
    r = client.post("/customers", json={"name": "Reschedule Me", "phone": "07700900070"}, headers=auth_headers)
    cust_a = r.json()["id"]
    r = client.post("/customers", json={"name": "Blocks Slot", "phone": "07700900071"}, headers=auth_headers)
    cust_b = r.json()["id"]

    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": cust_a, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00"},
        headers=auth_headers,
    )
    booking_id = r.json()["id"]
    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": cust_b, "starts_at": "2026-09-01T09:30:00", "ends_at": "2026-09-01T10:00:00"},
        headers=auth_headers,
    )
    assert r.status_code == 201, r.text

    # Re-import row A, moving it from 09:00 to the already-full 09:30 slot.
    up = upload(client, auth_headers, event["id"], [row_for("Reschedule Me", phone="07700900070", time="09:30", booking_id=booking_id)])
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["rows"][0]["status"] == "flagged"
    res = commit(client, auth_headers, event["id"], up["batch_id"])
    assert res.json()["batch"]["rows_flagged"] == 1
    original = db.get(Appointment, uuid.UUID(booking_id))
    assert original.starts_at.strftime("%H:%M") == "09:00"


def test_import_cancelled_active_transitions(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=1)

    # New row with status=cancelled creates a cancelled appointment and never touches stage.
    up = upload(client, auth_headers, event["id"], [row_for("Cancelled New", phone="07700900080", status="cancelled")])
    preview(client, auth_headers, event["id"], up["batch_id"])
    res = commit(client, auth_headers, event["id"], up["batch_id"])
    assert res.json()["batch"]["bookings_created"] == 1
    cust = db.query(Customer).filter(Customer.name == "Cancelled New").first()
    assert cust.stage != "appointment_booked"
    appt = db.query(Appointment).filter(Appointment.customer_id == cust.id).first()
    assert appt.status == "cancelled"

    # Fill the only slot with someone else.
    r = client.post("/customers", json={"name": "Slot Filler", "phone": "07700900081"}, headers=auth_headers)
    filler_id = r.json()["id"]
    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": filler_id, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00"},
        headers=auth_headers,
    )
    assert r.status_code == 201, r.text

    # Reactivating the cancelled booking into the now-full slot must be flagged, not silently applied.
    up2 = upload(client, auth_headers, event["id"], [row_for("Cancelled New", phone="07700900080", status="booked", booking_id=str(appt.id))])
    prev2 = preview(client, auth_headers, event["id"], up2["batch_id"])
    assert prev2["rows"][0]["status"] == "flagged"
    db.refresh(appt)
    assert appt.status == "cancelled"


# ---------------------------------------------------------------------------
# 14-16: isolation / provisional references
# ---------------------------------------------------------------------------


def test_import_rollback_after_partial_customer_write(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=1)
    rows = [
        row_for("Fills Slot", phone="07700900090", time="09:00"),
        row_for("Blocked New Customer", phone="07700900091", time="09:00"),
    ]
    up = upload(client, auth_headers, event["id"], rows)
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["rows"][0]["status"] == "created"
    assert prev["rows"][1]["status"] == "flagged"
    # The second row's customer must never have been left behind despite being created mid-savepoint
    # before its booking step discovered the slot was full.
    assert db.query(Customer).filter(Customer.name == "Blocked New Customer").count() == 0


def test_import_no_row_specific_audit_for_flagged_rows(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=1)
    rows = [
        row_for("Takes Slot", phone="07700900100", time="09:00"),
        row_for("Rejected Row", phone="07700900101", time="09:00"),
    ]
    up = upload(client, auth_headers, event["id"], rows)
    preview(client, auth_headers, event["id"], up["batch_id"])
    res = commit(client, auth_headers, event["id"], up["batch_id"])
    assert res.status_code == 200, res.text

    rejected = db.query(Customer).filter(Customer.name == "Rejected Row").first()
    assert rejected is None
    batch_id = uuid.UUID(up["batch_id"])
    logs_for_batch = db.query(AuditLog).filter(AuditLog.entity_id == batch_id).all()
    assert len(logs_for_batch) == 1
    assert logs_for_batch[0].action == "import.completed"
    for log in db.query(AuditLog).filter(AuditLog.action == "customer.import_created").all():
        if log.after:
            assert log.after.get("name") != "Rejected Row"


def test_import_preview_provisional_references_not_persisted(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=5)
    rows = [
        row_for("Provisional One", phone="07700900110", time="09:00"),
        row_for("Provisional One Dup", phone="+447700900110", time="09:00"),
    ]
    up = upload(client, auth_headers, event["id"], rows)
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["rows"][0]["customer_id"] is None
    assert prev["rows"][1]["customer_id"] is None
    assert prev["rows"][1]["mapped_data"]["provisional_customer_of_row"] == 1
    assert db.query(Customer).filter(Customer.name.like("Provisional One%")).count() == 0


# ---------------------------------------------------------------------------
# 17-19: concurrency and crash recovery
# ---------------------------------------------------------------------------


def test_import_concurrent_commits_separate_sessions(client, auth_headers, db, engine):
    event = create_event(client, auth_headers)
    up = upload(client, auth_headers, event["id"], [row_for("Concurrent Lead", phone="07700900120")])
    preview(client, auth_headers, event["id"], up["batch_id"])

    Session = sessionmaker(bind=engine)
    results = []
    errors = []

    def worker():
        session = Session()
        try:
            batch = session.get(ImportBatch, uuid.UUID(up["batch_id"]))
            evt = session.get(Event, uuid.UUID(event["id"]))
            user = session.query(User).filter(User.id == batch.owner_user_id).first()
            result = run_commit(session, batch, event=evt, user=user)
            results.append(result.status)
        except Exception as exc:  # pragma: no cover - diagnostic only
            errors.append(exc)
        finally:
            session.close()

    threads = [threading.Thread(target=worker) for _ in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    assert not errors, errors
    # The atomic claim guarantees only one thread actually performs the writes; the other may
    # observe "committing" (still in flight) or "committed" (already finished) depending on
    # scheduling, but never reprocesses. Poll briefly for the winner to finish.
    assert all(s in ("committing", "committed") for s in results)
    deadline = datetime.now(timezone.utc) + timedelta(seconds=5)
    final_batch = db.get(ImportBatch, uuid.UUID(up["batch_id"]))
    while final_batch.status != "committed" and datetime.now(timezone.utc) < deadline:
        db.expire(final_batch)
        final_batch = db.get(ImportBatch, uuid.UUID(up["batch_id"]))
    assert final_batch.status == "committed"
    assert db.query(Customer).filter(Customer.name == "Concurrent Lead").count() == 1
    assert db.query(Appointment).filter(Appointment.event_id == uuid.UUID(event["id"])).count() == 1


def test_import_concurrent_with_manual_booking_capacity(client, auth_headers, db, engine):
    event = create_event(client, auth_headers, slot_capacity=1)
    up = upload(client, auth_headers, event["id"], [row_for("Import Racer", phone="07700900130", time="09:00")])
    preview(client, auth_headers, event["id"], up["batch_id"])

    r = client.post("/customers", json={"name": "Manual Racer", "phone": "07700900131"}, headers=auth_headers)
    manual_customer_id = r.json()["id"]

    Session = sessionmaker(bind=engine)
    outcomes = []

    def import_worker():
        session = Session()
        try:
            batch = session.get(ImportBatch, uuid.UUID(up["batch_id"]))
            evt = session.get(Event, uuid.UUID(event["id"]))
            user = session.query(User).filter(User.id == batch.owner_user_id).first()
            result = run_commit(session, batch, event=evt, user=user)
            outcomes.append(("import", result.status, result.bookings_created))
        finally:
            session.close()

    def manual_worker():
        session = Session()
        try:
            evt = session.get(Event, uuid.UUID(event["id"]))
            user = session.query(User).filter(User.email == auth_headers_email[0]).first()
            payload = AppointmentCreate(customer_id=uuid.UUID(manual_customer_id), starts_at=datetime(2026, 9, 1, 9, 0), ends_at=datetime(2026, 9, 1, 9, 30))
            try:
                api_create_appointment(uuid.UUID(event["id"]), payload, session, user)
                outcomes.append(("manual", "ok", None))
            except Exception as exc:
                outcomes.append(("manual", "rejected", str(exc)))
        finally:
            session.close()

    auth_headers_email = [None]

    def find_owner_email():
        batch = db.get(ImportBatch, uuid.UUID(up["batch_id"]))
        owner = db.query(User).filter(User.id == batch.owner_user_id).first()
        auth_headers_email[0] = owner.email

    find_owner_email()

    threads = [threading.Thread(target=import_worker), threading.Thread(target=manual_worker)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    final_count = db.query(Appointment).filter(Appointment.event_id == uuid.UUID(event["id"]), Appointment.status != "cancelled").count()
    assert final_count == 1, outcomes


def test_import_slot_fills_between_preview_and_commit(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=1)
    up = upload(client, auth_headers, event["id"], [row_for("Late Importer", phone="07700900140", time="09:00")])
    prev = preview(client, auth_headers, event["id"], up["batch_id"])
    assert prev["rows"][0]["status"] == "created"

    r = client.post("/customers", json={"name": "Faster Booker", "phone": "07700900141"}, headers=auth_headers)
    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": r.json()["id"], "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00"},
        headers=auth_headers,
    )
    assert r.status_code == 201, r.text

    res = commit(client, auth_headers, event["id"], up["batch_id"])
    body = res.json()
    assert body["rows"][0]["status"] == "flagged"
    # The row's SAVEPOINT rollback must undo the provisional customer it created before
    # discovering the slot was gone, same as the pre-write-ambiguity case.
    assert db.query(Customer).filter(Customer.name == "Late Importer").count() == 0


def test_import_unauthorized_batch_access_rejected(client, auth_headers, db):
    event = create_event(client, auth_headers)
    up = upload(client, auth_headers, event["id"], [row_for("Private Lead", phone="07700900150")])

    other_headers, _ = register(client)
    r = client.get(f"/events/{event['id']}/imports/{up['batch_id']}", headers=other_headers)
    assert r.status_code == 404

    other_event = create_event(client, other_headers, name="Other User Event", day="2026-09-05")
    r = client.get(f"/events/{other_event['id']}/imports/{up['batch_id']}", headers=auth_headers)
    assert r.status_code == 404


def test_import_abandoned_committing_batch_recovery(client, auth_headers, db):
    event = create_event(client, auth_headers)
    up = upload(client, auth_headers, event["id"], [row_for("Recovered Lead", phone="07700900160")])
    preview(client, auth_headers, event["id"], up["batch_id"])

    batch = db.get(ImportBatch, uuid.UUID(up["batch_id"]))
    batch.status = "committing"
    batch.committing_started_at = datetime.now(timezone.utc) - timedelta(minutes=30)
    db.add(batch)
    db.commit()

    res = commit(client, auth_headers, event["id"], up["batch_id"])
    assert res.status_code == 200, res.text
    assert res.json()["batch"]["status"] == "committed"
    assert db.query(Customer).filter(Customer.name == "Recovered Lead").count() == 1


def test_import_full_reimport_idempotent(client, auth_headers, db):
    event = create_event(client, auth_headers)
    event_id = uuid.UUID(event["id"])
    up = upload(client, auth_headers, event["id"], [row_for("Idempotent Lead", phone="07700900170", email="idem@example.com")])
    preview(client, auth_headers, event["id"], up["batch_id"])
    commit(client, auth_headers, event["id"], up["batch_id"])

    assert db.query(Customer).filter(Customer.name == "Idempotent Lead").count() == 1
    assert db.query(Appointment).filter(Appointment.event_id == event_id).count() == 1

    r = client.get(f"/events/{event['id']}/export.xlsx", headers=auth_headers)
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    exported_rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]

    up2 = upload(client, auth_headers, event["id"], exported_rows)
    prev2 = preview(client, auth_headers, event["id"], up2["batch_id"])
    assert prev2["rows"][0]["status"] == "unchanged"
    res2 = commit(client, auth_headers, event["id"], up2["batch_id"])
    body2 = res2.json()
    assert body2["batch"]["customers_created"] == 0
    assert body2["batch"]["bookings_created"] == 0
    assert db.query(Customer).filter(Customer.name == "Idempotent Lead").count() == 1
    assert db.query(Appointment).filter(Appointment.event_id == event_id).count() == 1


# ---------------------------------------------------------------------------
# Merge
# ---------------------------------------------------------------------------


def test_merge_preserves_history_and_never_deletes_appointments(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=5)
    r = client.post("/customers", json={"name": "Survivor", "phone": "07700900200"}, headers=auth_headers)
    survivor_id = r.json()["id"]
    r = client.post("/customers", json={"name": "Loser", "email": "loser@example.com", "company": "Loser Co"}, headers=auth_headers)
    loser_id = r.json()["id"]

    # self-merge rejected
    r = client.post(f"/customers/{survivor_id}/merge/preview", json={"loser_customer_id": survivor_id}, headers=auth_headers)
    assert r.status_code == 400

    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": survivor_id, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00"},
        headers=auth_headers,
    )
    survivor_appt_id = r.json()["id"]
    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": loser_id, "starts_at": "2026-09-01T09:30:00", "ends_at": "2026-09-01T10:00:00", "notes": "loser booking"},
        headers=auth_headers,
    )
    loser_appt_id = r.json()["id"]

    r = client.post(f"/customers/{survivor_id}/merge/preview", json={"loser_customer_id": loser_id}, headers=auth_headers)
    assert r.status_code == 200, r.text
    analysis = r.json()
    assert len(analysis["divergent_bookings"]) == 0  # different times -> no_conflict, not divergent
    fingerprint = analysis["fingerprint"]

    r = client.post(
        f"/customers/{survivor_id}/merge/commit",
        json={"loser_customer_id": loser_id, "confirmed_fingerprint": fingerprint},
        headers=auth_headers,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["survivor"]["email"] == "loser@example.com"
    assert body["survivor"]["company"] == "Loser Co"

    assert db.get(Customer, uuid.UUID(loser_id)) is None
    # both appointments preserved, reassigned to survivor -- never deleted
    a1 = db.get(Appointment, uuid.UUID(survivor_appt_id))
    a2 = db.get(Appointment, uuid.UUID(loser_appt_id))
    assert a1 is not None and a2 is not None
    assert str(a1.customer_id) == survivor_id
    assert str(a2.customer_id) == survivor_id
    assert a2.notes == "loser booking"

    log = db.query(AuditLog).filter(AuditLog.action == "customer.merged").first()
    assert log is not None
    assert log.meta.get("loser_id") == loser_id
    # no fabricated Interaction should exist for either customer
    from app.db.models import Interaction
    assert db.query(Interaction).filter(Interaction.customer_id == uuid.UUID(survivor_id)).count() == 0


def test_merge_consolidates_identical_active_duplicate_booking(client, auth_headers, db):
    # Same event, same exact slot, same status/notes/staff on both sides -- an unambiguous
    # duplicate. The merge must not leave the survivor double-booked in that slot: one copy
    # stays active, the redundant one is cancelled (never deleted, so history survives).
    event = create_event(client, auth_headers, slot_capacity=5)
    r = client.post("/customers", json={"name": "Dup Survivor", "phone": "07700900220"}, headers=auth_headers)
    survivor_id = r.json()["id"]
    r = client.post("/customers", json={"name": "Dup Loser", "email": "duploser@example.com"}, headers=auth_headers)
    loser_id = r.json()["id"]

    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": survivor_id, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00"},
        headers=auth_headers,
    )
    survivor_appt_id = r.json()["id"]
    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": loser_id, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00"},
        headers=auth_headers,
    )
    loser_appt_id = r.json()["id"]

    r = client.post(f"/customers/{survivor_id}/merge/preview", json={"loser_customer_id": loser_id}, headers=auth_headers)
    analysis = r.json()
    assert len(analysis["identical_bookings"]) == 1
    assert analysis["identical_bookings"][0]["both_active"] is True
    assert len(analysis["capacity_conflicts"]) == 0  # identical, not divergent -- not blocking

    r = client.post(
        f"/customers/{survivor_id}/merge/commit",
        json={"loser_customer_id": loser_id, "confirmed_fingerprint": analysis["fingerprint"]},
        headers=auth_headers,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["cancelled_duplicate_appointment_ids"] == [loser_appt_id]

    survivor_appt = db.get(Appointment, uuid.UUID(survivor_appt_id))
    loser_appt = db.get(Appointment, uuid.UUID(loser_appt_id))
    assert survivor_appt.status == "booked"
    assert loser_appt.status == "cancelled"  # preserved, not deleted, but no longer occupying the slot
    assert str(loser_appt.customer_id) == survivor_id

    # The slot is not double-booked: only one active appointment remains at 09:00.
    active_at_slot = (
        db.query(Appointment)
        .filter(Appointment.event_id == uuid.UUID(event["id"]), Appointment.status != "cancelled")
        .count()
    )
    assert active_at_slot == 1


def test_merge_blocks_on_divergent_active_duplicate_booking(client, auth_headers, db):
    # Same event, same exact slot, but DIFFERENT status/notes -- ambiguous which one is right.
    # The merge must refuse rather than silently pick one or leave both active.
    event = create_event(client, auth_headers, slot_capacity=5)
    r = client.post("/customers", json={"name": "Conflict Survivor", "phone": "07700900230"}, headers=auth_headers)
    survivor_id = r.json()["id"]
    r = client.post("/customers", json={"name": "Conflict Loser", "email": "conflictloser@example.com"}, headers=auth_headers)
    loser_id = r.json()["id"]

    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": survivor_id, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00", "notes": "survivor notes"},
        headers=auth_headers,
    )
    survivor_appt_id = r.json()["id"]
    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": loser_id, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00", "notes": "loser notes -- different"},
        headers=auth_headers,
    )
    loser_appt_id = r.json()["id"]

    r = client.post(f"/customers/{survivor_id}/merge/preview", json={"loser_customer_id": loser_id}, headers=auth_headers)
    analysis = r.json()
    assert len(analysis["divergent_bookings"]) == 1
    assert len(analysis["capacity_conflicts"]) == 1  # both active + divergent -> blocking

    r = client.post(
        f"/customers/{survivor_id}/merge/commit",
        json={"loser_customer_id": loser_id, "confirmed_fingerprint": analysis["fingerprint"]},
        headers=auth_headers,
    )
    assert r.status_code == 409, r.text
    assert r.json()["detail"]["requires_manual_resolution"] is True

    # Nothing was committed: both customers and both active bookings still exist, untouched.
    assert db.get(Customer, uuid.UUID(loser_id)) is not None
    assert db.get(Appointment, uuid.UUID(survivor_appt_id)).status == "booked"
    assert db.get(Appointment, uuid.UUID(loser_appt_id)).status == "booked"
    assert str(db.get(Appointment, uuid.UUID(loser_appt_id)).customer_id) == loser_id  # not reassigned

    # Resolving by cancelling the loser's conflicting booking lets the merge proceed cleanly.
    r = client.patch(
        f"/events/{event['id']}/appointments/{loser_appt_id}",
        json={"status": "cancelled"},
        headers=auth_headers,
    )
    assert r.status_code == 200, r.text
    r = client.post(f"/customers/{survivor_id}/merge/preview", json={"loser_customer_id": loser_id}, headers=auth_headers)
    analysis2 = r.json()
    assert len(analysis2["capacity_conflicts"]) == 0
    r = client.post(
        f"/customers/{survivor_id}/merge/commit",
        json={"loser_customer_id": loser_id, "confirmed_fingerprint": analysis2["fingerprint"]},
        headers=auth_headers,
    )
    assert r.status_code == 200, r.text
    assert db.get(Customer, uuid.UUID(loser_id)) is None


def test_merge_changed_preview_requires_reconfirmation(client, auth_headers, db):
    r = client.post("/customers", json={"name": "Stable Survivor", "phone": "07700900210"}, headers=auth_headers)
    survivor_id = r.json()["id"]
    r = client.post("/customers", json={"name": "Drifting Loser"}, headers=auth_headers)
    loser_id = r.json()["id"]

    r = client.post(f"/customers/{survivor_id}/merge/preview", json={"loser_customer_id": loser_id}, headers=auth_headers)
    stale_fingerprint = r.json()["fingerprint"]

    # Something changes about the loser after the preview was taken.
    client.patch(f"/customers/{loser_id}", json={"email": "newly-added@example.com"}, headers=auth_headers)
    client.patch(f"/customers/{survivor_id}", json={"email": "conflicting@example.com"}, headers=auth_headers)

    r = client.post(
        f"/customers/{survivor_id}/merge/commit",
        json={"loser_customer_id": loser_id, "confirmed_fingerprint": stale_fingerprint},
        headers=auth_headers,
    )
    assert r.status_code == 409
    detail = r.json()["detail"]
    assert detail["requires_reconfirmation"] is True
    assert db.get(Customer, uuid.UUID(loser_id)) is not None  # nothing committed


# ---------------------------------------------------------------------------
# Regression: manual booking endpoints after appointment_rules.py extraction
# ---------------------------------------------------------------------------


def test_manual_booking_capacity_regression_after_extraction(client, auth_headers, db):
    event = create_event(client, auth_headers, slot_capacity=1)
    r = client.post("/customers", json={"name": "First Booker"}, headers=auth_headers)
    c1 = r.json()["id"]
    r = client.post("/customers", json={"name": "Second Booker"}, headers=auth_headers)
    c2 = r.json()["id"]

    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": c1, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00"},
        headers=auth_headers,
    )
    assert r.status_code == 201, r.text

    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": c2, "starts_at": "2026-09-01T09:00:00", "ends_at": "2026-09-01T09:30:00"},
        headers=auth_headers,
    )
    assert r.status_code == 409, r.text
    assert "full" in r.json()["detail"].lower()

    # misaligned duration still rejected with 400
    r = client.post(
        f"/events/{event['id']}/appointments",
        json={"customer_id": c2, "starts_at": "2026-09-01T10:00:00", "ends_at": "2026-09-01T10:15:00"},
        headers=auth_headers,
    )
    assert r.status_code == 400
